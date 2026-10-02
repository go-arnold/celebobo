import io
from datetime import UTC, datetime
from decimal import Decimal

from openpyxl import load_workbook

from apps.documents.domain.tables import Column, Table
from apps.documents.services.writers import CsvWriter, PdfTableWriter, XlsxWriter, writer_registry

TABLE = Table(
    title="Ventes",
    columns=(
        Column("product", "Produit"),
        Column("total", "Total", numeric=True),
        Column("at", "Date"),
    ),
    rows=(
        ('=HYPERLINK("x")', Decimal("12.50"), datetime(2026, 1, 5, 9, 30, tzinfo=UTC)),
        ("Coque", None, None),
    ),
)


def test_csv_is_excel_friendly_and_formula_safe():
    content = CsvWriter().write(TABLE).decode()

    assert content.startswith("﻿Produit;Total;Date")
    assert "'=HYPERLINK" in content
    assert "12.50;2026-01-05 10:30" in content


def test_xlsx_keeps_numbers_numeric():
    workbook = load_workbook(io.BytesIO(XlsxWriter().write(TABLE)))
    sheet = workbook.active

    assert sheet.title == "Ventes"
    assert [cell.value for cell in sheet[1]] == ["Produit", "Total", "Date"]
    assert sheet["B2"].value == 12.5
    assert sheet["A1"].font.bold


def test_pdf_tables_go_through_the_renderer(renderer):
    content = PdfTableWriter(renderer, {"company_name": "Celebobo"}).write(TABLE)

    template, context = renderer.calls[0]
    assert content == b"%PDF-fake"
    assert template == "documents/table.html"
    assert context["rows"][1] == ["Coque", "", ""]


def test_every_format_has_a_writer():
    assert list(writer_registry) == ["csv", "pdf", "xlsx"]
