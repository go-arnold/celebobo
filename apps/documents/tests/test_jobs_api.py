import io
from datetime import timedelta

import pytest
from django.utils import timezone
from openpyxl import load_workbook

from apps.accounts.tests.factories import ResellerFactory
from apps.documents.domain.enums import JobKind
from apps.documents.facades import DocumentJobFacade
from apps.documents.models import Job
from apps.documents.providers import register_generators
from apps.documents.services.generators import generator_registry
from core.container import container

pytestmark = pytest.mark.django_db

JOBS = "/api/v1/jobs/"


class Exploding:
    def run(self, request):
        raise RuntimeError("boom")


@pytest.fixture
def sold(reseller, phone, sell):
    return sell(reseller, phone)


@pytest.fixture
def exploding_export():
    generator_registry.add(JobKind.SALES_EXPORT.value, Exploding, replace=True)
    yield
    register_generators(container)


def export_sales(api, **payload):
    return api.post("/api/v1/bo/sales/export/", payload, format="json")


class TestSalesExport:
    def test_csv_export_runs_in_the_background(self, as_user, reseller, sold, run_jobs):
        api = as_user(reseller)

        with run_jobs():
            response = export_sales(api, format="csv", period="7d")

        assert response.status_code == 202
        assert response.json()["status"] == "pending"
        job = api.get(f"{JOBS}{response.json()['id']}/").json()
        assert job["status"] == "done"
        assert job["downloadable"] is True
        assert job["summary"] == {"rows": 1}
        assert job["filename"].startswith("ventes-")
        download = api.get(f"{JOBS}{job['id']}/download/")
        assert download["Content-Type"].startswith("text/csv")
        assert "attachment" in download["Content-Disposition"]
        body = download.content.decode()
        assert "Galaxy A55" in body
        assert "Jean Mukendi" in body

    def test_exports_respect_the_requester_scope(self, as_user, reseller, phone, sell, run_jobs):
        sell(reseller, phone)
        sell(ResellerFactory.create(), phone, sold_to="Autre client")
        api = as_user(reseller)

        with run_jobs():
            job_id = export_sales(api, format="xlsx").json()["id"]

        sheet = load_workbook(io.BytesIO(api.get(f"{JOBS}{job_id}/download/").content)).active
        assert sheet.max_row == 2
        assert sheet["L2"].value == "Jean Mukendi"

    def test_pdf_export(self, as_user, manager, sold, run_jobs, renderer):
        with run_jobs():
            job_id = export_sales(as_user(manager), format="pdf").json()["id"]

        template, context = renderer.calls[0]
        assert template == "documents/table.html"
        assert context["table"].totals == (("Chiffre d'affaires net", "100.00"),)
        assert Job.objects.get(pk=job_id).content_type == "application/pdf"

    def test_validates_the_format(self, as_user, reseller):
        assert export_sales(as_user(reseller), format="docx").status_code == 400

    def test_clients_cannot_export(self, as_user, shopper):
        assert export_sales(as_user(shopper), format="csv").status_code == 403


class TestOtherJobs:
    def test_products_export(self, as_user, manager, phone, run_jobs):
        api = as_user(manager)
        with run_jobs():
            job_id = api.post("/api/v1/bo/products/export/", {"format": "csv"}).json()["id"]

        lines = api.get(f"{JOBS}{job_id}/download/").content.decode().splitlines()
        assert lines[0].lstrip("\ufeff").split(";") == [
            "slug",
            "name",
            "category",
            "price",
            "sale_price",
            "cost_price",
            "stock",
            "stock_threshold",
            "is_active",
            "sales_count",
        ]
        assert lines[1].startswith("galaxy-a55;Galaxy A55;smartphones;120.00;;80.00;20")
        assert api.post("/api/v1/bo/products/export/", {"format": "pdf"}).status_code == 400

    def test_analytics_report(self, as_user, reseller, sold, run_jobs, renderer):
        with run_jobs():
            response = as_user(reseller).post("/api/v1/bo/analytics/export/", {"period": "30d"})

        template, context = renderer.calls[0]
        assert response.status_code == 202
        assert template == "documents/analytics_report.html"
        assert context["sellers"] == []
        assert Job.objects.get().filename.startswith("rapport-")


class TestLifecycle:
    def test_jobs_are_private(self, as_user, reseller, run_jobs):
        with run_jobs():
            job_id = export_sales(as_user(reseller), format="csv").json()["id"]
        other = as_user(ResellerFactory.create())

        assert other.get(f"{JOBS}{job_id}/").status_code == 404
        assert other.get(f"{JOBS}{job_id}/download/").status_code == 404
        assert other.get(JOBS).json()["meta"]["count"] == 0

    def test_listing(self, as_user, reseller, run_jobs):
        api = as_user(reseller)
        with run_jobs():
            export_sales(api, format="csv")

        body = api.get(JOBS).json()
        assert body["meta"]["count"] == 1
        assert body["results"][0]["kind"] == "sales_export"

    def test_pending_jobs_cannot_be_downloaded(self, as_user, reseller):
        api = as_user(reseller)
        job_id = export_sales(api, format="csv").json()["id"]

        response = api.get(f"{JOBS}{job_id}/download/")

        assert response.status_code == 409
        assert response.json()["code"] == "job_not_ready"

    @pytest.mark.usefixtures("exploding_export")
    def test_crashes_mark_the_job_failed(self, as_user, reseller, run_jobs):
        with run_jobs():
            export_sales(as_user(reseller), format="csv")

        job = Job.objects.get()
        assert job.status == "failed"
        assert job.error.startswith("Une erreur interne")

    def test_purge_removes_old_jobs(self, as_user, reseller, run_jobs):
        with run_jobs():
            export_sales(as_user(reseller), format="csv")
        Job.objects.update(created_at=timezone.now() - timedelta(days=30))

        assert container.resolve(DocumentJobFacade).purge() == 1
        assert not Job.objects.exists()


@pytest.mark.usefixtures("sold")
@pytest.mark.parametrize(
    ("url", "payload"),
    [
        ("/api/v1/bo/sales/export/", {"format": "pdf"}),
        ("/api/v1/bo/analytics/export/", {"period": "7d"}),
    ],
)
def test_pdf_templates_render(as_user, manager, run_jobs, url, payload):
    api = as_user(manager)
    with run_jobs():
        job_id = api.post(url, payload, format="json").json()["id"]

    assert api.get(f"{JOBS}{job_id}/download/").content.startswith(b"%PDF")
